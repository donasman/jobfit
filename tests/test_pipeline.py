from pathlib import Path

from openpyxl import load_workbook

from jobfit import pipeline
from jobfit.excel.workbook import SHEET_JOBS, SHEET_QUESTIONS, create_jobs_workbook
from jobfit.paths import Project
from jobfit.state import State

FIXTURES = Path(__file__).parent / "fixtures"


def _project(tmp_path, rows, questions=()):
    project = Project(tmp_path)
    create_jobs_workbook(project.jobs_xlsx)
    wb = load_workbook(project.jobs_xlsx)
    # 수식 열이 500행까지 차 있어 append는 그 뒤에 붙는다 — 2행부터 직접 쓴다
    for i, r in enumerate(rows):
        for col, v in enumerate(r, start=1):
            wb[SHEET_JOBS].cell(row=2 + i, column=col, value=v)
    for i, q in enumerate(questions):
        for col, v in enumerate(q, start=1):
            wb[SHEET_QUESTIONS].cell(row=2 + i, column=col, value=v)
    wb.save(project.jobs_xlsx)
    project.state_dir.mkdir()
    project.output_dir.mkdir()
    return project


def test_sync_in_registers_and_removes(tmp_path):
    project = _project(tmp_path, [["예시랩", "백엔드", "https://www.saramin.co.kr/zf_user/jobs/view?rec_idx=10000001", None, "", None, None]])
    state = State()
    report = pipeline.sync_in(project, state)
    assert [j.key for j in report.added] == ["saramin:10000001"]
    state.save(project.state_file)

    # 행을 지우면 state에서만 빠진다
    wb = load_workbook(project.jobs_xlsx)
    wb[SHEET_JOBS].delete_rows(2)
    wb.save(project.jobs_xlsx)
    state = State.load(project.state_file)
    report = pipeline.sync_in(project, state)
    assert [j.key for j in report.removed] == ["saramin:10000001"] and not state.jobs


def test_duplicate_rows_same_key_are_skipped(tmp_path):
    project = _project(
        tmp_path,
        [
            ["A", "", "https://www.saramin.co.kr/zf_user/jobs/view?rec_idx=1", None, "", None, None],
            ["A", "", "https://m.saramin.co.kr/zf_user/jobs/relay/view?rec_idx=1&utm_source=mail", None, "", None, None],
        ],
    )
    report = pipeline.sync_in(project, State())
    assert len(report.added) == 1 and len(report.duplicates) == 1


def test_task_runs_once_per_change(tmp_path):
    project = _project(tmp_path, [["A", "", "https://www.wanted.co.kr/wd/1", None, "", None, "재실행"]])
    state = State()
    r1 = pipeline.sync_in(project, state)
    assert r1.tasks == []  # 새 공고의 첫 값은 그냥 기록만 한다
    r2 = pipeline.sync_in(project, state)
    assert r2.tasks == []  # 값이 안 바뀌면 다시 걸리지 않는다
    wb = load_workbook(project.jobs_xlsx)
    wb[SHEET_JOBS].cell(row=2, column=7).value = None  # cell(value=None)은 아무것도 안 바꾼다
    wb.save(project.jobs_xlsx)
    r3 = pipeline.sync_in(project, state)
    assert r3.tasks == [] and state.get("wanted:1").task_seen == ""
    wb[SHEET_JOBS].cell(row=2, column=7).value = "재실행"
    wb.save(project.jobs_xlsx)
    r4 = pipeline.sync_in(project, state)
    assert [(j.company, t) for j, t in r4.tasks] == [("A", "재실행")]
    r5 = pipeline.sync_in(project, state)
    assert r5.tasks == []  # 처리된 뒤 같은 값이 남아 있어도 다시 걸리지 않는다


def test_questions_match_by_company_or_company_title(tmp_path):
    project = _project(
        tmp_path,
        [
            ["예시컴퍼니", "Node.js개발자", "https://www.jobkorea.co.kr/Recruit/GI_Read/10000002", None, "", None, None],
            ["예시랩", "백엔드", "https://www.saramin.co.kr/zf_user/jobs/view?rec_idx=10000001", None, "", None, None],
            ["예시랩", "프런트", "https://www.saramin.co.kr/zf_user/jobs/view?rec_idx=99", None, "", None, None],
        ],
        [["예시컴퍼니", 1, "지원동기", None, None], ["예시랩 / 백엔드", 1, "협업", None, None], ["예시랩", 2, "모호", None, None]],
    )
    report = pipeline.sync_in(project, State())
    assert set(report.questions) == {"jobkorea:10000002", "saramin:10000001"}
    assert [q.text for q in report.unmatched_questions] == ["모호"]


def test_clean_from_fixture_without_network(tmp_path):
    project = _project(tmp_path, [["예시랩", "백엔드", "https://www.saramin.co.kr/zf_user/jobs/view?rec_idx=10000001", None, "", None, None]])
    state = State()
    pipeline.sync_in(project, state)
    job = state.get("saramin:10000001")
    job_dir = project.job_dir(job.key)
    job_dir.mkdir(parents=True)
    (job_dir / "source.html").write_text((FIXTURES / "saramin-sample.html").read_text(encoding="utf-8", errors="replace"), encoding="utf-8")
    (job_dir / "source-detail.html").write_text((FIXTURES / "saramin-sample-detail.html").read_text(encoding="utf-8", errors="replace"), encoding="utf-8")
    (job_dir / "fetch.json").write_text('{"ok": true, "title": "예시랩 백엔드 (Backend) 개발자"}', encoding="utf-8")
    job.mark("fetch", done=True)
    ok, msg = pipeline._STEP_RUNNERS["clean"](project, job, job_dir)
    assert ok, msg
    posting = (job_dir / "posting.txt").read_text(encoding="utf-8")
    assert len(posting) <= 2500
    assert "자격" in posting or "업무" in posting
    assert (job_dir / "meta.json").exists() and (job_dir / "posting.full.txt").exists()


def test_fetch_reuses_original_already_fetched_by_add(tmp_path):
    url = "https://www.saramin.co.kr/zf_user/jobs/view?rec_idx=10000001"
    project = _project(tmp_path, [["예시랩", "백엔드", url, None, "", None, None]])
    state = State()
    pipeline.sync_in(project, state)
    job = state.get("saramin:10000001")
    job_dir = project.job_dir(job.key)
    job_dir.mkdir(parents=True)
    (job_dir / "source.html").write_text("<html></html>", encoding="utf-8")
    (job_dir / "fetch.json").write_text(f'{{"ok": true, "url": "{url}", "title": "t"}}', encoding="utf-8")
    ok, _ = pipeline._STEP_RUNNERS["fetch"](project, job, job_dir)  # 네트워크 없이 통과해야 한다
    assert ok and job.steps["fetch"]["note"] == "이미 받은 원문 사용"
    pipeline.forget_fetch(job_dir)
    assert not (job_dir / "fetch.json").exists()


def test_user_posting_wins_over_fetched_page_without_body(tmp_path):
    from jobfit.clean.posting import clean_job

    project = _project(tmp_path, [["원티드사", "서버", "https://www.wanted.co.kr/wd/1", None, "", None, None]])
    job_dir = project.job_dir("wanted:1")
    job_dir.mkdir(parents=True)
    (job_dir / "source.html").write_text("<html><body><div>로그인</div><div>회원가입</div></body></html>", encoding="utf-8")
    (job_dir / "fetch.json").write_text('{"ok": true, "url": "https://www.wanted.co.kr/wd/1"}', encoding="utf-8")
    r1 = clean_job(job_dir, project.data_dir)
    assert not r1.ok and "본문이 없습니다" in r1.error
    pasted = "주요업무\n- API 개발\n자격요건\n- Java 3년 이상\n문의 a@b.com"
    (job_dir / "posting.txt").write_text(pasted, encoding="utf-8")
    r2 = clean_job(job_dir, project.data_dir)
    assert r2.ok and r2.source == "pasted"
    r3 = clean_job(job_dir, project.data_dir)  # 다시 돌려도 붙여 넣은 글을 계속 쓰고, 고치지 않는다
    assert r3.source == "pasted"
    assert (job_dir / "posting.txt").read_text(encoding="utf-8") == pasted


def test_machine_posting_is_rewritten_from_source(tmp_path):
    from jobfit.clean.posting import clean_job

    project = _project(tmp_path, [["예시랩", "백엔드", "https://www.saramin.co.kr/zf_user/jobs/view?rec_idx=10000001", None, "", None, None]])
    job_dir = project.job_dir("saramin:10000001")
    job_dir.mkdir(parents=True)
    (job_dir / "source.html").write_text((FIXTURES / "saramin-sample.html").read_text(encoding="utf-8", errors="replace"), encoding="utf-8")
    (job_dir / "source-detail.html").write_text((FIXTURES / "saramin-sample-detail.html").read_text(encoding="utf-8", errors="replace"), encoding="utf-8")
    (job_dir / "fetch.json").write_text('{"ok": true, "title": "t"}', encoding="utf-8")
    assert clean_job(job_dir, project.data_dir).source == "fetched"
    assert clean_job(job_dir, project.data_dir).source == "fetched"  # 스크립트가 쓴 posting.txt는 사용자 글로 오인하지 않는다


def test_pasted_posting_is_used_when_no_url(tmp_path):
    project = _project(tmp_path, [["수기회사", "서버", "", "2026-12-01", "", None, None]])
    state = State()
    pipeline.sync_in(project, state)
    job = next(iter(state.jobs.values()))
    lines = pipeline.run(project, state, job)
    assert "실패" in lines[0] and job.status == "오류"
    job_dir = project.job_dir(job.key)
    (job_dir / "posting.txt").write_text("담당업무\n- API 개발\n자격요건\n- Java 3년\n우대사항\n- AWS", encoding="utf-8")
    job.reset_from("fetch")
    lines = pipeline.run(project, state, job)
    assert job.step_done("fetch") and job.step_done("clean"), lines
    assert job.status != "오류"


# --- 첫 코스: 공고 링크 하나 → report 한 장 ---


def _fresh(tmp_path) -> Project:
    from jobfit import cli

    assert cli.main(["--root", str(tmp_path), "init"]) == 0
    return Project(tmp_path)


def test_next_without_jobs_leads_with_a_job_link(tmp_path):
    # 처음 쓰는 사람의 첫 코스 — 내 자료보다 공고 링크 하나가 먼저다
    project = _fresh(tmp_path)
    action = pipeline.next_action(project, State())
    assert action.command == 'jobfit add "<공고 링크>"'
    assert "먼저 공고 링크 하나를 넣고 결과를 보세요" in action.message
    assert "경험 카드가 없어도 공고 요건 정리(report.md)까지 나옵니다" in action.message
    assert "경험 카드를 넣으면 매칭 점수와 이력서 강조 포인트가" in action.message


def test_next_with_sources_but_no_jobs_still_shows_the_job_link(tmp_path):
    # sources/ 에 자료만 넣은 사람에게도 공고 한 장 길을 알려 준다(자료 안내는 그대로)
    project = _fresh(tmp_path)
    (project.sources_dir / "r.txt").write_text("주문 API 개선 프로젝트", encoding="utf-8")
    action = pipeline.next_action(project, State())
    assert action.command == "jobfit exp extract"
    assert pipeline.FIRST_JOB_TIP in action.message


def test_before_the_first_report_the_job_comes_before_my_files(tmp_path, capsys):
    from jobfit import cli
    from jobfit.state import STATUS_AI, Job

    project = _fresh(tmp_path)
    (project.sources_dir / "r.txt").write_text("주문 API 개선 프로젝트", encoding="utf-8")
    assert cli.main(["--root", str(tmp_path), "exp", "extract"]) == 0
    capsys.readouterr()
    state = State()
    job = Job(key="saramin:1", site="saramin", company="예시랩", row=2, status=STATUS_AI)
    job.mark("fetch")
    job.mark("clean")
    state.jobs[job.key] = job

    # 둘 다 AI 단계 — report 한 장이 나오기 전에는 공고가 먼저
    action = pipeline.next_action(project, state)
    assert action.job is job and "AI 단계(analyze)" in action.message
    # report 가 한 번 나온 뒤에는 매칭에 쓰일 경험 카드가 먼저
    job.mark("analyze")
    job.mark("match")
    job.mark("report")
    assert "경험 카드 뽑기" in pipeline.next_action(project, state).message
