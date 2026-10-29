from datetime import datetime

from openpyxl import load_workbook

from jobfit.excel.reader import read_jobs_workbook
from jobfit.excel.workbook import JOB_HEADERS, QUESTION_HEADERS, SHEET_JOBS, SHEET_QUESTIONS, create_jobs_workbook


def _fill(ws, rows, start=2):
    # 공고 시트는 수식 열이 500행까지 차 있어 append가 그 뒤에 붙는다 — 사용자처럼 2행부터 쓴다
    for i, values in enumerate(rows):
        for col, v in enumerate(values, start=1):
            ws.cell(row=start + i, column=col, value=v)


def _make(tmp_path, job_rows, question_rows=()):
    path = tmp_path / "jobs.xlsx"
    create_jobs_workbook(path)
    wb = load_workbook(path)
    _fill(wb[SHEET_JOBS], job_rows)
    _fill(wb[SHEET_QUESTIONS], question_rows)
    wb.save(path)
    return path


def test_headers_and_sheets_exist(tmp_path):
    path = _make(tmp_path, [])
    wb = load_workbook(path)
    assert [c.value for c in wb[SHEET_JOBS][1]][: len(JOB_HEADERS)] == list(JOB_HEADERS)
    assert [c.value for c in wb[SHEET_QUESTIONS][1]] == list(QUESTION_HEADERS)
    assert "안내" in wb.sheetnames


def test_reads_rows_and_parses_dates(tmp_path):
    path = _make(
        tmp_path,
        [
            ["예시랩", "백엔드 개발자", "https://www.saramin.co.kr/zf_user/jobs/relay/view?rec_idx=10000001", datetime(2026, 10, 27), "", None, None],
            ["예시컴퍼니", "Node.js개발자(신입)", "https://www.jobkorea.co.kr/Recruit/GI_Read/10000002", "2026.11.15", "메일 접수", "보류", ""],
            ["예시페이", "Server Developer", "", "2026년 11월 30일", "", "지원완료", "재실행"],
            ["상시회사", "", "", "채용시 마감", "", None, None],
        ],
        [["예시랩 / 백엔드 개발자", 1, "지원 동기를 기술하시오", "800", "제외"], ["예시컴퍼니", "", "협업 경험", None, None]],
    )
    result = read_jobs_workbook(path)
    assert result.ok, [str(e) for e in result.errors]
    deadlines = [j.deadline for j in result.jobs]
    assert deadlines[:3] == ["2026-10-27", "2026-11-15", "2026-11-30"]
    assert result.jobs[3].deadline == "" and result.jobs[3].deadline_note == "채용시 마감"
    assert result.jobs[0].decision == "미정" and result.jobs[1].decision == "보류"
    assert result.jobs[2].task == "재실행"
    assert result.questions[0].limit == 800 and result.questions[0].count_spaces is False
    assert result.questions[1].number == 0 and result.questions[1].count_spaces is True


def test_hyperlink_with_display_text_uses_target(tmp_path):
    path = _make(tmp_path, [["회사", "", "공고 보기", None, "", None, None]])
    wb = load_workbook(path)
    cell = wb[SHEET_JOBS].cell(row=2, column=3)
    cell.hyperlink = "https://www.saramin.co.kr/zf_user/jobs/view?rec_idx=42"
    wb.save(path)
    result = read_jobs_workbook(path)
    assert result.ok and result.jobs[0].url.endswith("rec_idx=42")


def test_bad_rows_are_reported_and_others_survive(tmp_path):
    path = _make(
        tmp_path,
        [
            ["", "직무만", "", None, "", None, None],
            ["회사", "", "ftp://nope", None, "", None, None],
            ["회사2", "", "", "11월 어느 날", "", None, None],
            ["회사3", "", "", None, "", "합격", None],
            ["정상", "", "https://www.wanted.co.kr/wd/1", None, "", None, None],
        ],
    )
    result = read_jobs_workbook(path)
    assert len(result.errors) == 4
    assert [j.company for j in result.jobs] == ["정상"]
    rows = sorted(e.row for e in result.errors)
    assert rows == [2, 3, 4, 5]


def test_column_order_change_is_fine(tmp_path):
    path = _make(tmp_path, [])
    wb = load_workbook(path)
    ws = wb[SHEET_JOBS]
    for col, name in enumerate(["메모", "링크", "회사명", "직무", "마감일", "결정", "작업"], start=1):
        ws.cell(row=1, column=col, value=name)
    _fill(ws, [["m", "https://www.wanted.co.kr/wd/7", "원티드사", "서버", None, None, None]])
    wb.save(path)
    result = read_jobs_workbook(path)
    assert result.ok, [str(e) for e in result.errors]
    assert result.jobs[0].company == "원티드사" and result.jobs[0].url.endswith("/wd/7")


def test_missing_required_header_stops_sheet(tmp_path):
    path = _make(tmp_path, [])
    wb = load_workbook(path)
    wb[SHEET_JOBS].cell(row=1, column=1, value="기업")
    wb.save(path)
    result = read_jobs_workbook(path)
    assert any("필수 헤더" in e.message and "회사명" in e.message for e in result.errors)
    assert result.jobs == []
