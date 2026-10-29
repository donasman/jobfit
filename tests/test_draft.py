import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

from jobfit import cli
from jobfit.ai.draft import parse_answers
from jobfit.excel.workbook import SHEET_JOBS, SHEET_QUESTIONS
from jobfit.paths import Project
from jobfit.state import State

FIXTURES = Path(__file__).parent / "fixtures"
URL = "https://www.saramin.co.kr/zf_user/jobs/view?rec_idx=10000001"
KEY = "saramin:10000001"
ANALYSIS = {
    "role_summary": "제조 AI 솔루션 백엔드 개발",
    "must": ["Python 기반 API 개발", "SQL 성능 최적화"],
    "nice": ["Docker 경험"],
    "duties": ["백엔드 설계 및 개발"],
    "values": ["주도성"],
    "keywords": ["Python", "Docker"],
    "questions_found": [],
}
CARDS = """
- id: api
  title: 주문 조회 API 응답 시간 개선
  period: 2025-03 ~ 2025-05
  role: 백엔드 개발
  tags: [Python, 성능 최적화, Docker, SQL]
  summary: N+1 제거와 캐시
  star:
    situation: 주문 목록 화면이 느리다는 문의가 반복됐다.
    task: 기존 API 계약을 바꾸지 않고 응답 시간을 줄인다.
    action: 쿼리 로그로 N+1을 찾고 fetch join과 캐시를 붙였다.
    result: 응답 시간이 줄고 쿼리 수가 고정됐다.
  facts: ["p95 1.8초 → 0.4초", "주문 1건당 DB 쿼리 21회 → 1회"]
  usable_for: [직무역량, 문제해결]
  verified: true
- id: doc
  title: 온보딩 문서 정리
  tags: [협업, 문서화]
  summary: 세팅 문서
  star: {situation: s, task: t, action: a, result: r}
  facts: ["세팅 소요 3일 → 1일"]
  usable_for: [협업, 성장]
  verified: false
"""
GOOD_Q1 = "주문 목록 화면이 느리다는 문의가 반복됐습니다. 쿼리 로그를 확인해 N+1 문제를 찾았고, 주문 1건당 21회였던 DB 쿼리를 1회로 줄였습니다. 그 결과 p95가 1.8초에서 0.4초가 됐습니다."
GOOD_Q2 = "새 팀원이 혼자 세팅하기 어려웠습니다. 막힌 지점을 기록해 문서를 정리했고 세팅 소요가 3일에서 1일로 줄었습니다. [확인 필요: 문서를 본 인원]"


def _project(tmp_path) -> Project:
    project = Project(tmp_path)
    assert cli.main(["--root", str(tmp_path), "init"]) == 0
    assert cli.main(["--root", str(tmp_path), "add", URL, "--company", "예시랩", "--title", "백엔드 개발자", "--deadline", "2026-10-27"]) == 0
    job_dir = project.job_dir(KEY)
    job_dir.mkdir(parents=True)
    shutil.copy(FIXTURES / "saramin-sample.html", job_dir / "source.html")
    shutil.copy(FIXTURES / "saramin-sample-detail.html", job_dir / "source-detail.html")
    (job_dir / "fetch.json").write_text(json.dumps({"ok": True, "url": URL, "title": "t", "files": {}}), encoding="utf-8")
    project.experiences_file.write_text(CARDS, encoding="utf-8")
    wb = load_workbook(project.jobs_xlsx)
    ws = wb[SHEET_QUESTIONS]
    for r, (num, text, limit) in enumerate([(1, "어려운 문제를 해결한 경험을 적어주세요", len(GOOD_Q1)), (2, "동료와 협업한 경험을 적어주세요", len(GOOD_Q2))], start=2):
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = "예시랩", num, text, limit
    wb.save(project.jobs_xlsx)
    assert cli.main(["--root", str(tmp_path), "run"]) == 0
    (project.ai_dir(KEY) / "analyze.response.json").write_text(json.dumps(ANALYSIS, ensure_ascii=False), encoding="utf-8")
    assert cli.main(["--root", str(tmp_path), "accept", "analyze", "--job", KEY]) == 0
    assert cli.main(["--root", str(tmp_path), "run"]) == 0  # 매칭·report → 초안 AI 단계
    return project


def _job(project):
    return State.load(project.state_file).get(KEY)


def _answer(tmp_path, text, name="a.txt"):
    p = tmp_path / name
    p.write_text(text, encoding="utf-8")
    return str(p)


def test_draft_prompt_contents(tmp_path):
    project = _project(tmp_path)
    job = _job(project)
    assert job.status == "AI대기" and job.steps["draft"]["phase"] == "draft"
    prompt = (project.ai_dir(KEY) / "draft.prompt.md").read_text(encoding="utf-8")
    assert prompt.startswith("[jobfit · 자소서 초안]") and "draft.response.md" in prompt.splitlines()[0]
    assert "=== q1 ===" in prompt and "=== q2 ===" in prompt
    assert '쓸 수 있는 사실: "p95 1.8초 → 0.4초"; "주문 1건당 DB 쿼리 21회 → 1회"' in prompt
    assert "  - 행동: 쿼리 로그로 N+1을 찾고" in prompt
    assert f"글자수: {len(GOOD_Q1)}자 (공백 포함)" in prompt
    assert "사실 확인 전 카드" in prompt  # doc 카드 verified false
    assert "{{" not in prompt and len(prompt) <= 8000


def test_good_draft_finishes_with_report_and_dashboard(tmp_path, capsys):
    project = _project(tmp_path)
    capsys.readouterr()
    answer = f"초안입니다.\n=== q1 ===\n{GOOD_Q1}\n=== q2 ===\n{GOOD_Q2}\n"
    assert cli.main(["--root", str(tmp_path), "accept", "--file", _answer(tmp_path, answer)]) == 0
    out = capsys.readouterr().out
    assert "초안을 받아들였습니다" in out and "확인 필요 1" in out
    job = _job(project)
    assert job.step_done("draft") and job.status == "검토대기"
    job_dir = project.job_dir(KEY)
    assert (job_dir / "drafts" / "q1.md").read_text(encoding="utf-8").strip() == GOOD_Q1
    lint = json.loads((job_dir / "lint.json").read_text(encoding="utf-8"))
    assert [q["errors"] for q in lint["questions"]] == [0, 0]
    report = (job_dir / "report.md").read_text(encoding="utf-8")
    assert "### 1. 어려운 문제를 해결한 경험을 적어주세요" in report and GOOD_Q1 in report
    assert "초안의 [확인 필요] 1곳" in report and "내 말투로 다시" in report
    cli.main(["--root", str(tmp_path), "run"])
    ws = load_workbook(project.dashboard_xlsx)["검토"]
    header = [c.value for c in ws[1]]
    row = {h: ws.cell(row=2, column=i + 1).value for i, h in enumerate(header)}
    # q2 가 사실 확인 전 카드(doc)의 숫자 3일·1일을 써서 경고 2건
    assert row["초안"] == "경고 2건" and row["처리상태"] == "검토대기"
    assert "사실 확인 전(verified: false) 카드의 숫자" in report


def test_fabricated_number_goes_to_revise_once(tmp_path, capsys):
    project = _project(tmp_path)
    bad_q1 = GOOD_Q1.replace("0.4초가 됐습니다.", "0.4초로 78% 빨라졌습니다.")[: len(GOOD_Q1)]
    assert cli.main(["--root", str(tmp_path), "accept", "--file", _answer(tmp_path, f"=== q1 ===\n{bad_q1}\n=== q2 ===\n{GOOD_Q2}")]) == 0
    out = capsys.readouterr().out
    assert "1개 문항만 한 번 고쳐" in out and "'78%'" in out
    job = _job(project)
    assert job.status == "AI대기" and job.steps["draft"]["phase"] == "revise" and job.steps["draft"]["revise_qids"] == ["q1"]
    revise = (project.ai_dir(KEY) / "revise.prompt.md").read_text(encoding="utf-8")
    assert "=== q1 ===" in revise and "=== q2 ===" not in revise and "'78%'은(는) 경험 카드 facts에 없는 숫자" in revise
    assert cli.main(["--root", str(tmp_path), "prep"]) == 0  # 이제 수정 프롬프트를 준다
    assert "revise.prompt.md" in capsys.readouterr().out
    assert cli.main(["--root", str(tmp_path), "accept", "--file", _answer(tmp_path, f"=== q1 ===\n{GOOD_Q1}", "r.txt")]) == 0
    assert "수정본을 받아들였습니다" in capsys.readouterr().out
    job = _job(project)
    assert job.step_done("draft") and job.steps["draft"]["revised"] is True and job.steps["draft"]["errors"] == 0
    assert (project.job_dir(KEY) / "drafts" / "q2.md").read_text(encoding="utf-8").strip() == GOOD_Q2  # 걸리지 않은 문항은 그대로


def test_revise_that_still_fails_ends_with_marked_errors(tmp_path, capsys):
    project = _project(tmp_path)
    bad = f"=== q1 ===\n{GOOD_Q1[:-5]} 매출 5억 원.\n=== q2 ===\n{GOOD_Q2}"
    cli.main(["--root", str(tmp_path), "accept", "--file", _answer(tmp_path, bad)])
    cli.main(["--root", str(tmp_path), "accept", "--file", _answer(tmp_path, bad.split("=== q2")[0], "r.txt")])
    job = _job(project)
    assert job.step_done("draft") and job.steps["draft"]["errors"] >= 1
    report = (project.job_dir(KEY) / "report.md").read_text(encoding="utf-8")
    assert "⛔" in report and "고치지 못한 오류" in report
    ws = load_workbook(project.dashboard_xlsx)["검토"]
    cli.main(["--root", str(tmp_path), "dashboard"])
    ws = load_workbook(project.dashboard_xlsx)["검토"]
    header = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=header.index("초안") + 1).value.startswith("오류")


def test_unparseable_answer_makes_retry(tmp_path, capsys):
    project = _project(tmp_path)
    assert cli.main(["--root", str(tmp_path), "accept", "--file", _answer(tmp_path, "초안을 써 드릴게요! 첫 문항은...")]) == 1
    assert "구분선을 찾지 못했습니다" in capsys.readouterr().out
    retry = (project.ai_dir(KEY) / "draft.retry.prompt.md").read_text(encoding="utf-8")
    assert retry.startswith("[jobfit · 자소서 초안 다시 요청 1]") and "## 앞 답의 문제" in retry
    assert _job(project).steps["draft"]["prompt"] == "draft.retry.prompt.md"


def test_cache_and_rewrite_task(tmp_path, capsys):
    project = _project(tmp_path)
    cli.main(["--root", str(tmp_path), "accept", "--file", _answer(tmp_path, f"=== q1 ===\n{GOOD_Q1}\n=== q2 ===\n{GOOD_Q2}")])
    # 매칭 입력이 그대로인 채 draft만 지우면 캐시로 AI 없이 다시 채운다
    state = State.load(project.state_file)
    state.get(KEY).reset_from("draft")
    state.save(project.state_file)
    capsys.readouterr()
    cli.main(["--root", str(tmp_path), "run"])
    assert "초안까지 완료" in capsys.readouterr().out and _job(project).steps["draft"]["cached"] is True
    # '초안만 재작성' 작업은 캐시를 쓰지 않는다
    wb = load_workbook(project.jobs_xlsx)
    wb[SHEET_JOBS]["G2"] = "초안만 재작성"
    wb.save(project.jobs_xlsx)
    cli.main(["--root", str(tmp_path), "run"])
    job = _job(project)
    assert job.status == "AI대기" and "draft" in job.force_steps
    cli.main(["--root", str(tmp_path), "accept", "--file", _answer(tmp_path, f"=== q1 ===\n{GOOD_Q1}\n=== q2 ===\n{GOOD_Q2}", "n.txt")])
    assert "draft" not in _job(project).force_steps


def test_parse_answers_edge_cases():
    answers, errors, warnings = parse_answers("```\n=== Q1 ===\n가\n=== q3 ===\n다\n```", ["q1", "q2"])
    assert answers == {"q1": "가"} and errors == ["q2 본문이 없습니다"] and "q3" in warnings[0]


def test_exp_check_warns_star_numbers_outside_facts(tmp_path, capsys):
    cli.main(["--root", str(tmp_path), "init"])
    project = Project(tmp_path)
    project.experiences_file.write_text(CARDS.replace("응답 시간이 줄고", "응답 시간이 78% 줄고"), encoding="utf-8")
    cli.main(["--root", str(tmp_path), "exp", "check"])
    assert "star에 facts에 없는 숫자가 있습니다: 78%" in capsys.readouterr().out


def test_question_block_shrinks_but_keeps_facts():
    from jobfit.ai.draft import _question_block

    long_star = {k: "가" * 200 for k in ("situation", "task", "action", "result")}
    cards = {
        "a": {"id": "a", "title": "주력", "star": long_star, "facts": ["p95 1.8초 → 0.4초"], "verified": True},
        "b": {"id": "b", "title": "보조", "star": long_star, "facts": ["쿼리 21회 → 1회"], "verified": True},
    }
    q = {"id": "q1", "text": "문항", "limit": 500, "count_spaces": True, "type": "직무역량", "cards": [{"id": "a", "role": "주력"}, {"id": "b", "role": "보조"}]}
    full, one, two = (_question_block(q, cards, level) for level in (0, 1, 2))
    assert len(full) > len(one) > len(two)
    assert "가" * 200 in full and "가" * 120 + "…" in one and "가" * 60 + "…" in two
    assert two.count("가" * 60) == 4  # 보조 카드 STAR는 빠지고 주력 카드 네 칸만
    for block in (full, one, two):
        assert '"p95 1.8초 → 0.4초"' in block and '"쿼리 21회 → 1회"' in block  # facts는 끝까지 남는다
