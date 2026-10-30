import warnings

import pytest
from openpyxl import load_workbook

from jobfit import cli
from jobfit.excel.append import ExcelBusy, append_job_row
from jobfit.excel.reader import read_jobs_workbook
from jobfit.excel.workbook import SHEET_JOBS, SHEET_QUESTIONS, create_jobs_workbook
from jobfit.paths import Project


def _project(tmp_path):
    project = Project(tmp_path)
    assert cli.main(["--root", str(tmp_path), "init"]) == 0
    return project


def test_append_fills_first_empty_row_and_keeps_user_rows(tmp_path):
    path = tmp_path / "jobs.xlsx"
    create_jobs_workbook(path)
    wb = load_workbook(path)
    wb[SHEET_JOBS]["A2"] = "기존회사"
    wb[SHEET_JOBS]["F2"] = "보류"
    wb.save(path)

    row = append_job_row(path, tmp_path / "backups", "예시컴퍼니", "Node.js개발자(신입)", "https://www.jobkorea.co.kr/Recruit/GI_Read/10000002", "2026-11-15", "메모")
    assert row == 3
    result = read_jobs_workbook(path)
    assert result.ok, [str(e) for e in result.errors]
    assert [(j.company, j.decision) for j in result.jobs] == [("기존회사", "보류"), ("예시컴퍼니", "미정")]
    assert result.jobs[1].deadline == "2026-11-15" and result.jobs[1].url.endswith("10000002")
    assert len(list((tmp_path / "backups").glob("jobs_*.xlsx"))) == 1


def test_append_keeps_validations_and_formula(tmp_path):
    path = tmp_path / "jobs.xlsx"
    create_jobs_workbook(path)
    before = {s: len(load_workbook(path)[s].data_validations.dataValidation) for s in (SHEET_JOBS, SHEET_QUESTIONS)}
    append_job_row(path, tmp_path / "b", "A")
    append_job_row(path, tmp_path / "b", "B")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(path)
    after = {s: len(wb[s].data_validations.dataValidation) for s in (SHEET_JOBS, SHEET_QUESTIONS)}
    assert before == after and after[SHEET_JOBS] == 4 and after[SHEET_QUESTIONS] == 3  # 결정·작업·마감일·지원방법
    assert str(wb[SHEET_JOBS]["I3"].value).startswith("=IF(")  # 공고표시 수식이 그대로
    assert wb[SHEET_JOBS].column_dimensions["I"].hidden


def test_append_refuses_when_excel_has_file_open(tmp_path):
    path = tmp_path / "jobs.xlsx"
    create_jobs_workbook(path)
    (tmp_path / "~$jobs.xlsx").write_bytes(b"lock")
    before = path.read_bytes()
    with pytest.raises(ExcelBusy):
        append_job_row(path, tmp_path / "b", "A")
    assert path.read_bytes() == before


def test_cli_add_with_all_fields_needs_no_network(tmp_path, capsys):
    project = _project(tmp_path)
    url = "https://www.saramin.co.kr/zf_user/jobs/relay/view?rec_idx=10000001"
    rc = cli.main(["--root", str(tmp_path), "add", url, "--company", "예시랩", "--title", "백엔드 개발자", "--deadline", "2026.10.27"])
    assert rc == 0
    out = capsys.readouterr().out
    assert "공고 시트 2행" in out and "saramin:10000001" in out
    # 같은 공고(다른 꼴의 링크)를 다시 넣으면 추가하지 않는다
    rc = cli.main(["--root", str(tmp_path), "add", "https://m.saramin.co.kr/zf_user/jobs/view?rec_idx=10000001&utm_source=x", "--company", "x", "--title", "y", "--deadline", "2026-10-27"])
    assert rc == 0 and "이미 공고 시트 2행" in capsys.readouterr().out
    assert [j.company for j in read_jobs_workbook(project.jobs_xlsx).jobs] == ["예시랩"]


def test_cli_add_apply_via(tmp_path, capsys):
    project = _project(tmp_path)
    assert cli.main(["--root", str(tmp_path), "add", "--company", "수기", "--apply-via", "이메일"]) == 0
    assert read_jobs_workbook(project.jobs_xlsx).jobs[0].apply_via == "이메일"
    assert "열이 없어" not in capsys.readouterr().out


def test_cli_add_apply_via_on_old_workbook_says_column_is_missing(tmp_path, capsys):
    project = _project(tmp_path)
    wb = load_workbook(project.jobs_xlsx)
    wb[SHEET_JOBS].delete_cols(8)  # 지원방법 열이 없던 예전 jobs.xlsx
    wb.save(project.jobs_xlsx)
    assert cli.main(["--root", str(tmp_path), "add", "--company", "수기", "--apply-via", "홈페이지"]) == 0
    assert "'지원방법' 열이 없어" in capsys.readouterr().out
    result = read_jobs_workbook(project.jobs_xlsx)
    assert result.ok and result.jobs[0].company == "수기" and result.jobs[0].apply_via == ""


def test_cli_add_without_url_and_bad_inputs(tmp_path, capsys):
    _project(tmp_path)
    assert cli.main(["--root", str(tmp_path), "add"]) == 2
    assert cli.main(["--root", str(tmp_path), "add", "ftp://x"]) == 2
    assert cli.main(["--root", str(tmp_path), "add", "--company", "수기", "--deadline", "11월쯤"]) == 2
    assert cli.main(["--root", str(tmp_path), "add", "--company", "수기", "--title", "서버"]) == 0
    assert cli.main(["--root", str(tmp_path), "add", "--company", "수기", "--title", "서버"]) == 0
    assert "이미 공고 시트 2행" in capsys.readouterr().out
