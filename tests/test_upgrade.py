"""예전 버전 폴더에 새 zip을 덮어써도 안 깨지는지 — 전방 호환 회귀.

배포는 기존 폴더 위에 덮어쓰는 방식이라 사용자 데이터(jobs.xlsx · me/ · sources/ · output/ · .jobfit/)는 예전 모양 그대로 남는다.
여기서는 v0.1~v0.2 시절 모양의 작업 폴더를 그 자리에서 만들고, 지금 코드가 오류 없이 읽고
없는 것은 "없어서 못 하는 일"만 조용히 안내하는지 본다. 바이너리 파일은 저장소에 넣지 않는다(jobs.xlsx는 openpyxl로).

예전 폴더의 모양
  - jobs.xlsx: v0.3.0에서 생긴 '지원방법' 열이 없다
  - me/profile.yaml: v0.2.8에서 생긴 파일이라 없다
  - sources/자소서/: v0.3.2에서 생긴 자리라 없다 — 원안 자소서는 예전 자리 me/samples/에 있다
  - me/experiences.yaml: 예전 형식(verified·source 없음, usable_for 가 예전 이름, tags 가 문자열 하나)
  - .jobfit/state.json: 예전 키만(apply_via·force_steps·user_questions 없음, 단계 기록에 해시 없음)
"""

from __future__ import annotations

import json

from openpyxl import load_workbook

from jobfit import cli, jobkey, profile, style
from jobfit.excel.append import append_job_row, append_question_rows
from jobfit.excel.reader import read_jobs_workbook
from jobfit.excel.workbook import JOB_HEADERS, SHEET_JOBS, SHEET_QUESTIONS, create_jobs_workbook
from jobfit.experiences import card_types, check_cards, file_sha, load_cards
from jobfit.match import inputs_sha, run_match
from jobfit.paths import Project
from jobfit.state import Job, State

OLD_CARDS = """# 경험 은행 — 예전 형식
- id: exp-old-api
  title: 주문 API 개선
  period: 2025.03~2025.05
  org: 예시 커머스
  role: 백엔드
  tags: [Java, Spring Boot]
  summary: "N+1 을 없애고 캐시를 붙여 응답 시간을 줄인 경험"
  star:
    situation: 주문 목록이 느렸다.
    task: 응답 시간을 줄인다.
    action: fetch join 으로 묶고 캐시를 붙였다.
    result: 응답 시간이 줄었다.
  facts:
    - "p95 1.8초 → 0.4초"
  usable_for: [기술역량, 문제해결]

- id: exp-old-doc
  title: 온보딩 문서 정리
  period: 2025-07
  tags: 문서화
  summary: "흩어진 실행 방법을 한 문서로 모은 경험"
  star:
    situation: 새 팀원이 며칠을 썼다.
    task: 첫날 실행까지 되는 문서를 만든다.
    action: 따라 하며 막힌 곳을 기록했다.
    result: 세팅이 빨라졌다.
  facts:
    - "문서 1개로 통합"
  usable_for: [팀워크]
"""

OLD_RESUME = """items:
  - section: 프로젝트
    title: 주문 API 개선
    text: "N+1 제거와 캐시로 p95 1.8초 → 0.4초"
    tags: [Java, Spring Boot]
"""

SAMPLE = "저는 문제를 먼저 정의합니다. 그래서 로그부터 봤습니다.\n\n그 결과 원인을 찾았습니다. 이후 팀에 공유했습니다.\n"
RULES = '# 문체 규칙\n## 쓰는 방식\n- 결론부터\n## 피하는 표현\n- "열심히" — 추상적\n## 대표 문단\n문단입니다.\n'
POSTING = "주요업무\n- 주문 API 개발과 성능 개선\n자격요건\n- Java 3년 이상\n- Spring Boot 경험\n우대사항\n- Redis 캐시 경험\n"
ANALYSIS = {
    "must": ["Java 3년 이상", "Spring Boot 경험"],
    "nice": ["Redis 캐시 경험"],
    "keywords": ["Java", "Spring Boot", "Redis"],
    "questions_found": [],
    "role_summary": "주문 API 개발",
}
COMPANY, TITLE, DEADLINE = "예전회사", "백엔드", "2026-12-01"


def old_jobs_xlsx(project: Project) -> None:
    """v0.3.0 이전 jobs.xlsx — '지원방법' 열이 없다. 공고 1건·문항 1건."""
    create_jobs_workbook(project.jobs_xlsx)
    wb = load_workbook(project.jobs_xlsx)
    ws = wb[SHEET_JOBS]
    ws.delete_cols(JOB_HEADERS.index("지원방법") + 1)
    for col, value in enumerate([COMPANY, TITLE, "", DEADLINE, "예전 메모", None, None], start=1):
        ws.cell(row=2, column=col, value=value)
    for col, value in enumerate([COMPANY, 1, "지원 동기를 적으시오", 800, None], start=1):
        wb[SHEET_QUESTIONS].cell(row=2, column=col, value=value)
    wb.save(project.jobs_xlsx)


def old_project(tmp_path) -> tuple[Project, str]:
    """예전 버전이 남긴 작업 폴더 전체. (프로젝트, 공고키)."""
    project = Project(tmp_path)
    old_jobs_xlsx(project)

    project.me_dir.mkdir()
    project.experiences_file.write_text(OLD_CARDS, encoding="utf-8")
    project.resume_file.write_text(OLD_RESUME, encoding="utf-8")
    project.legacy_samples_dir.mkdir()  # 예전 자리 me/samples — sources/자소서 는 없다
    (project.legacy_samples_dir / "s1.txt").write_text(SAMPLE, encoding="utf-8")
    me_style = style.style_dir(project.me_dir)
    me_style.mkdir()
    (me_style / "style_rules.md").write_text(RULES, encoding="utf-8")
    # 예전 수치 파일에는 samples_sha 가 없다
    (me_style / "style_stats.json").write_text(json.dumps({"samples": 1, "sentences": 4}, ensure_ascii=False), encoding="utf-8")

    project.sources_dir.mkdir()
    (project.sources_dir / "이력서.txt").write_text("주문 API 개선 프로젝트. Java, Spring Boot 를 썼다.\n", encoding="utf-8")

    key = jobkey.derive("", COMPANY, TITLE, DEADLINE).key
    job_dir = project.job_dir(key)
    job_dir.mkdir(parents=True)
    (job_dir / "posting.txt").write_text(POSTING, encoding="utf-8")
    (job_dir / "posting.full.txt").write_text(POSTING, encoding="utf-8")
    (job_dir / "analysis.json").write_text(json.dumps(ANALYSIS, ensure_ascii=False), encoding="utf-8")
    (job_dir / "report.md").write_text("# 예전 리포트\n", encoding="utf-8")

    project.state_dir.mkdir()
    project.state_file.write_text(json.dumps(old_state(key), ensure_ascii=False, indent=2), encoding="utf-8")
    # 예전 sources.json — cards·direct 같은 나중에 생긴 칸이 없다
    (project.state_dir / "sources.json").write_text(
        json.dumps({"files": {"sources/이력서.txt": {"sha1": "0" * 40, "slug": "이력서", "status": "done", "parts": 1}}}, ensure_ascii=False),
        encoding="utf-8",
    )
    return project, key


def old_state(key: str) -> dict:
    """v0.1~v0.2 state.json — apply_via·force_steps·user_questions 가 없고 단계 기록에 입력 해시도 없다."""
    at = "2026-01-02 03:04:05"
    return {
        "version": 1,
        "last_sync": at,
        "sheet_errors": [],
        "jobs": {
            key: {
                "key": key, "site": "local", "company": COMPANY, "title": TITLE, "url": "",
                "deadline": DEADLINE, "memo": "예전 메모", "decision": "미정", "task_seen": "",
                "status": "문항대기",
                "steps": {s: {"done": True, "at": at} for s in ("fetch", "clean", "analyze", "match", "report")},
                "error": "", "next_hint": "문항 시트에 문항 입력", "row": 2,
                "registered_at": at, "updated_at": at,
            }
        },
    }


# --- jobs.xlsx: 나중에 더한 열 ---


def test_old_jobs_xlsx_without_apply_via_reads_without_errors(tmp_path):
    project, _ = old_project(tmp_path)
    result = read_jobs_workbook(project.jobs_xlsx)
    assert result.ok, [str(e) for e in result.errors]  # 나중에 더한 열이 없는 건 오류가 아니다
    assert "지원방법" not in result.job_headers
    assert [(j.company, j.apply_via, j.deadline) for j in result.jobs] == [(COMPANY, "", DEADLINE)]
    assert [(q.job_ref, q.limit, q.count_spaces) for q in result.questions] == [(COMPANY, 800, True)]


def test_old_jobs_xlsx_still_takes_new_rows_and_questions(tmp_path, capsys):
    project, _ = old_project(tmp_path)
    # 사용자 명령으로 들어오는 길 — 열이 없다는 것만 알리고 나머지는 넣는다
    assert cli.main(["--root", str(tmp_path), "add", "", "--company", "보탠회사", "--apply-via", "홈페이지"]) == 0
    assert "'지원방법' 열이 없어" in capsys.readouterr().out
    assert append_job_row(project.jobs_xlsx, project.backups_dir, "새회사", "서버", "", "", "", "홈페이지") == 4
    assert append_question_rows(project.jobs_xlsx, project.backups_dir, COMPANY, [("협업 경험", 500, False)], 2) == [3]

    result = read_jobs_workbook(project.jobs_xlsx)
    assert result.ok, [str(e) for e in result.errors]
    # 열이 없으니 지원방법은 쓰지 않는다. 나머지는 그대로 들어간다
    assert [(j.company, j.apply_via) for j in result.jobs] == [(COMPANY, ""), ("보탠회사", ""), ("새회사", "")]
    assert [(q.number, q.text, q.limit, q.count_spaces) for q in result.questions][1] == (2, "협업 경험", 500, False)


# --- .jobfit/state.json: 나중에 생긴 칸 ---


def test_old_state_json_loads_with_later_fields_defaulted(tmp_path):
    project, key = old_project(tmp_path)
    state = State.load(project.state_file)
    job = state.get(key)
    assert job is not None and job.company == COMPANY
    assert job.apply_via == "" and job.force_steps == [] and job.user_questions == []
    assert job.step_done("report") and not job.step_done("draft")
    state.save(project.state_file)  # 다시 써도 그대로 읽힌다
    assert State.load(project.state_file).get(key).apply_via == ""


def test_state_json_with_unknown_or_odd_values_falls_back_to_defaults(tmp_path):
    path = tmp_path / "state.json"
    path.write_text(
        json.dumps({
            "version": 9, "last_sync": None, "sheet_errors": None,
            "jobs": {
                "a:1": {"key": "a:1", "site": "a", "company": "A", "나중칸": 1, "steps": None, "force_steps": None},
                "a:2": {"company": "B"},  # key·site 가 없는 줄
                "a:3": "망가진 줄",
            },
        }, ensure_ascii=False),
        encoding="utf-8",
    )
    state = State.load(path)
    assert sorted(state.jobs) == ["a:1", "a:2"]  # 모양이 아닌 줄만 버린다
    assert state.sheet_errors == [] and state.last_sync == ""
    assert state.jobs["a:1"].steps == {} and state.jobs["a:1"].force_steps == []
    assert state.jobs["a:2"].key == "a:2" and state.jobs["a:2"].site == "a"
    path.write_text("null", encoding="utf-8")
    assert State.load(path).jobs == {}


# --- me/: 자리를 옮긴 원안, 아직 없는 profile.yaml, 예전 형식 카드 ---


def test_originals_are_still_read_from_legacy_me_samples(tmp_path):
    project, _ = old_project(tmp_path)
    assert not project.samples_dir.exists()  # v0.3.2 에서 생긴 자리는 예전 폴더에 없다
    assert [p.name for p in style.list_samples(project)] == ["s1.txt"]
    assert style.draft_blocker(project) == ""  # 원안·문체 규칙이 다 있으니 초안을 쓸 수 있다
    assert not style.style_pending(project)
    # 같은 자소서를 새 자리에 복사해 두어도 한 번만 센다
    project.samples_dir.mkdir(parents=True)
    (project.samples_dir / "s1.txt").write_text(SAMPLE, encoding="utf-8")
    assert len(style.list_samples(project)) == 1


def test_old_style_stats_without_sha_is_recomputed_quietly(tmp_path):
    project, _ = old_project(tmp_path)
    stats = style.refresh_stats(project)
    assert stats is not None and stats["samples_sha"]  # 예전 파일에는 해시가 없어 다시 잰다
    assert style.refresh_stats(project) is None  # 두 번째는 다시 재지 않는다
    (style.style_dir(project.me_dir) / "style_stats.json").write_text("[1, 2]", encoding="utf-8")
    assert style.load_stats(project.me_dir) is None  # 모양이 다른 파일은 없는 것으로
    assert style.stats_lines(style.load_stats(project.me_dir)) == []


def test_missing_profile_yaml_is_quiet_guidance_not_an_error(tmp_path, capsys):
    project, key = old_project(tmp_path)
    assert not project.profile_file.exists()
    assert profile.load_profile(project.profile_file) == ([], [])
    assert cli.main(["--root", str(tmp_path), "profile", "check"]) == 0
    assert "me/profile.yaml 에 항목이 없습니다" in capsys.readouterr().out
    assert cli.main(["--root", str(tmp_path), "apply", "sheet", "--job", key]) == 0
    out = capsys.readouterr().out
    assert "me/profile.yaml" in out  # 없어서 못 하는 일만 알린다
    assert (project.apply_dir(key) / "fill.md").exists()


def test_old_experiences_yaml_reads_with_warnings_only(tmp_path):
    from jobfit.clean.keywords import KeywordDict

    project, _ = old_project(tmp_path)
    cards, problems = load_cards(project.experiences_file)
    assert problems == [] and [c["id"] for c in cards] == ["exp-old-api", "exp-old-doc"]
    issues = check_cards(cards, KeywordDict.load(project.data_dir))
    assert [i for i in issues if i.level == "오류"] == []  # 예전 형식은 경고까지만
    assert card_types(cards[0]) == ["직무역량", "문제해결"]  # 예전 이름도 알아듣는다
    assert card_types(cards[1]) == ["협업"]


# --- 폴더 전체 ---


def test_cli_runs_through_an_old_project_without_errors(tmp_path, capsys):
    project, key = old_project(tmp_path)
    for argv in (["check"], ["sync-in"], ["next"], ["run"], ["dashboard"], ["exp", "check"], ["style", "stats"]):
        assert cli.main(["--root", str(tmp_path), *argv]) == 0, argv
        capsys.readouterr()
    assert project.dashboard_xlsx.exists()
    report = (project.job_dir(key) / "report.md").read_text(encoding="utf-8")
    assert COMPANY in report and "## 매칭" in report  # 예전 report.md 를 새 형식으로 다시 썼다
    job = State.load(project.state_file).get(key)
    assert job.status != "오류" and job.error == ""
    assert job.user_questions and job.step_done("match")  # 문항 시트의 문항을 읽어 매칭을 다시 했다


# --- 끝난 일은 업그레이드해도 끝난 채로 ---

FINISHED_CARD = """- id: exp-a
  title: 주문 API 개선
  period: 2025-03 ~ 2025-05
  tags: [Java, Spring Boot]
  summary: "응답 시간을 줄인 경험"
  star: {situation: 느렸다., task: 줄인다., action: 캐시를 붙였다., result: 빨라졌다.}
  facts: ["p95 1.8초 → 0.4초"]
  usable_for: [지원동기, 직무역량]
  verified: true
"""
OLD_DRAFT = "예전 버전이 써 둔 초안입니다. 주문 API 를 개선했습니다."


def finished_old_project(tmp_path, *, rules_sha: bool) -> tuple[Project, str]:
    """초안까지 끝낸 예전 폴더. rules_sha=False 면 그 칸이 생기기 전(v0.2.5 이하)에 만든 기록이다."""
    project = Project(tmp_path)
    old_jobs_xlsx(project)
    project.me_dir.mkdir()
    project.experiences_file.write_text(FINISHED_CARD, encoding="utf-8")
    key = jobkey.derive("", COMPANY, TITLE, DEADLINE).key
    job_dir = project.job_dir(key)
    (job_dir / "drafts").mkdir(parents=True)
    (job_dir / "posting.txt").write_text(POSTING, encoding="utf-8")
    (job_dir / "posting.full.txt").write_text(POSTING, encoding="utf-8")
    (job_dir / "analysis.json").write_text(json.dumps(ANALYSIS, ensure_ascii=False), encoding="utf-8")
    (job_dir / "drafts" / "q1.md").write_text(OLD_DRAFT, encoding="utf-8")
    (job_dir / "lint.json").write_text(
        json.dumps({"questions": [{"qid": "q1", "chars": 40, "errors": 0, "warnings": 0, "confirm_needed": 0, "issues": []}]}, ensure_ascii=False),
        encoding="utf-8",
    )
    project.state_dir.mkdir()
    job = Job(key=key, site="local", company=COMPANY, title=TITLE, deadline=DEADLINE, row=2,
              user_questions=[{"text": "지원 동기를 적으시오", "limit": 800, "count_spaces": True, "row": 2}])
    run_match(project, job)
    shas = inputs_sha(project, job)
    if not rules_sha:
        shas.pop("rules_sha")  # v0.2.6에서 생긴 칸 — 그 전 기록에는 없다
    for step in ("fetch", "clean", "analyze"):
        job.mark(step, done=True)
    job.mark("match", done=True, score=70, questions=1, **shas)
    job.mark("report", done=True, resume_sha=file_sha(project.resume_file))
    job.mark("draft", done=True)
    job.status = "검토대기"
    (job_dir / "report.md").write_text(f"# 예전 리포트\n\n{OLD_DRAFT}\n", encoding="utf-8")
    State(jobs={key: job}).save(project.state_file)
    return project, key


def test_finished_old_job_is_not_redone_on_upgrade(tmp_path, capsys):
    project, key = finished_old_project(tmp_path, rules_sha=True)
    assert cli.main(["--root", str(tmp_path), "run"]) == 0
    capsys.readouterr()
    job = State.load(project.state_file).get(key)
    assert job.step_done("draft") and job.status == "검토대기"
    assert OLD_DRAFT in (project.job_dir(key) / "report.md").read_text(encoding="utf-8")


def test_finished_draft_survives_upgrade_from_before_rules_sha(tmp_path, capsys):
    project, key = finished_old_project(tmp_path, rules_sha=False)
    assert cli.main(["--root", str(tmp_path), "run"]) == 0
    capsys.readouterr()
    assert (project.job_dir(key) / "drafts" / "q1.md").exists()  # 파일은 남아 있다
    assert State.load(project.state_file).get(key).step_done("draft")
    assert OLD_DRAFT in (project.job_dir(key) / "report.md").read_text(encoding="utf-8")


# --- 반대 방향(왕복) ---


def test_data_written_by_this_version_reads_back(tmp_path):
    project = Project(tmp_path)
    create_jobs_workbook(project.jobs_xlsx)
    append_job_row(project.jobs_xlsx, project.backups_dir, "새회사", "서버", "https://www.wanted.co.kr/wd/1", "2026-12-24", "", "홈페이지")
    append_question_rows(project.jobs_xlsx, project.backups_dir, "새회사 / 서버", [("지원 동기", 800, True)], 1)
    result = read_jobs_workbook(project.jobs_xlsx)
    assert result.ok, [str(e) for e in result.errors]
    assert [(j.company, j.apply_via, j.deadline) for j in result.jobs] == [("새회사", "홈페이지", "2026-12-24")]
    assert [(q.job_ref, q.number, q.limit) for q in result.questions] == [("새회사 / 서버", 1, 800)]

    project.state_dir.mkdir(exist_ok=True)
    state = State.load(project.state_file)
    from jobfit import pipeline

    pipeline.sync_in(project, state)
    state.save(project.state_file)
    again = State.load(project.state_file)
    job = next(iter(again.jobs.values()))
    assert job.apply_via == "홈페이지" and [q["text"] for q in job.user_questions] == ["지원 동기"]

    project.me_dir.mkdir(exist_ok=True)
    profile.append_items(project.profile_file, project.backups_dir, [{"section": "학력", "school": "○○대학교", "verified": True}], "sources/이력서.txt")
    items, problems = profile.load_profile(project.profile_file)
    assert problems == [] and profile.item_name(items[0]) == "○○대학교"
    assert [i for i in profile.check_profile(items) if i.level == "오류"] == []

    project.samples_dir.mkdir(parents=True)
    (project.samples_dir / "s1.txt").write_text(SAMPLE, encoding="utf-8")
    written, unreadable = style.write_stats(project)
    assert unreadable == [] and style.load_stats(project.me_dir) == written
    assert style.refresh_stats(project) is None  # 방금 쓴 수치는 다시 재지 않는다
