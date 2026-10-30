import json
import shutil
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from jobfit import cli
from jobfit.excel.workbook import SHEET_QUESTIONS, header_letters
from jobfit.notices import HEADER
from jobfit.paths import Project
from jobfit.report import build_report
from jobfit.state import State

FIXTURES = Path(__file__).parent / "fixtures"
URL = "https://www.saramin.co.kr/zf_user/jobs/view?rec_idx=10000001"
KEY = "saramin:10000001"
ANALYSIS = {
    "role_summary": "제조 AI 솔루션 백엔드 개발",
    "must": ["Python 기반 API 개발", "Docker 경험", "SQL 성능 최적화"],
    "nice": ["Kotlin·Spring Boot 개발"],
    "duties": ["백엔드 설계 및 개발"],
    "values": [],
    "keywords": ["Python", "Docker", "Spring Boot", "MongoDB"],
    "questions_found": [],
}
CARDS = """
- id: etl
  title: ETL 파이프라인
  tags: [Python, PostgreSQL]
  summary: 배치 재처리
  facts: ["3천만 건"]
  usable_for: [직무역량, 문제해결]
  verified: true
- id: api
  title: 주문 API 개선 | 캐시
  tags: [Java, Spring Boot, Docker, 성능 최적화]
  summary: N+1 제거
  facts: ["p95 1.8초 → 0.4초"]
  usable_for: [문제해결, 협업]
  verified: false
"""
RESUME = """
items:
  - section: 프로젝트
    title: 주문 API 개선
    text: 스프링부트와 도커로 배포
    tags: [스프링부트, Docker]
  - section: 기술
    title: 언어
    text: 파이썬, 자바
    tags: [Python, Java]
"""


def _project(tmp_path, cards=CARDS, resume=None, analysis=ANALYSIS) -> Project:
    """사람인 픽스처 + 분석 답까지 넣은 프로젝트. 네트워크·AI 없이."""
    project = Project(tmp_path)
    assert cli.main(["--root", str(tmp_path), "init"]) == 0
    assert cli.main(["--root", str(tmp_path), "add", URL, "--company", "예시랩", "--title", "백엔드 개발자", "--deadline", "2026-10-27"]) == 0
    job_dir = project.job_dir(KEY)
    job_dir.mkdir(parents=True)
    shutil.copy(FIXTURES / "saramin-sample.html", job_dir / "source.html")
    shutil.copy(FIXTURES / "saramin-sample-detail.html", job_dir / "source-detail.html")
    (job_dir / "fetch.json").write_text(json.dumps({"ok": True, "url": URL, "title": "t", "files": {}}), encoding="utf-8")
    if cards is not None:
        project.experiences_file.write_text(cards, encoding="utf-8")
    if resume is not None:
        project.resume_file.write_text(resume, encoding="utf-8")
    assert cli.main(["--root", str(tmp_path), "run"]) == 0
    (project.ai_dir(KEY) / "analyze.response.json").write_text(json.dumps(analysis, ensure_ascii=False), encoding="utf-8")
    assert cli.main(["--root", str(tmp_path), "accept", "analyze", "--job", KEY]) == 0
    assert cli.main(["--root", str(tmp_path), "run"]) == 0
    return project


def _job(project):
    return State.load(project.state_file).get(KEY)


def test_no_questions_ends_at_report_with_need_questions(tmp_path, capsys):
    project = _project(tmp_path)
    job = _job(project)
    assert job.step_done("match") and job.step_done("report") and not job.step_done("draft")
    assert job.status == "문항대기"
    match = json.loads((project.job_dir(KEY) / "match.json").read_text(encoding="utf-8"))
    assert match["score"] is not None and "gaps" not in match
    capsys.readouterr()
    cli.main(["--root", str(tmp_path), "next"])
    out = capsys.readouterr().out
    # 사람은 메뉴나 문항 시트(공고 칸에 무엇을 적는지까지), 에이전트는 사용자가 준 원문을 question add 로
    assert "start.bat 2번(문항 넣기)나 jobs.xlsx 문항 시트(공고 칸에 '예시랩')에 넣고 jobfit run" in out
    assert f"jobfit question add --job {KEY}" in out


def test_report_sections(tmp_path):
    project = _project(tmp_path)
    text = build_report(project, _job(project), today=date(2026, 10, 20))
    assert text.startswith("# 예시랩 — 백엔드 개발자")
    assert "마감 2026-10-27 (D-7, 2026-10-20 기준)" in text
    for heading in ("## 공고 요약", "## 매칭", "## 이력서 강조 포인트", "## 자기소개서 문항", "## 확인 필요"):
        assert heading in text
    assert "부족" not in text and "갭" not in text
    assert "주문 API 개선 \\| 캐시 ⚠" in text  # 표 안의 | 이스케이프, verified false 표시
    assert "앞세울 경험 (내 이력서 항목이 없어 경험 카드 순위로)" in text
    assert "공고에 자소서 문항이 없습니다" in text
    assert "verified가 true가 아닙니다" in text
    assert "> 지원 방법 **홈페이지** · 근거: 공고 글 '접수방법 : 자유양식의 이력서를 홈페이지 접수'" in text


def test_user_apply_via_overrides_page_and_rebuilds_report(tmp_path, capsys):
    project = _project(tmp_path)
    wb = load_workbook(project.jobs_xlsx)
    ws = wb["공고"]
    ws[f"{header_letters(ws)['지원방법']}2"] = "이메일"
    wb.save(project.jobs_xlsx)
    capsys.readouterr()
    assert cli.main(["--root", str(tmp_path), "run"]) == 0
    assert "지원방법이 바뀌어 report부터 다시" in capsys.readouterr().out
    text = (project.job_dir(KEY) / "report.md").read_text(encoding="utf-8")
    assert "> 지원 방법 **이메일** · 근거: jobs.xlsx 지원방법 칸" in text
    ws = load_workbook(project.dashboard_xlsx)["검토"]
    assert ws.cell(row=2, column=[c.value for c in ws[1]].index("지원방법") + 1).value == "이메일"
    capsys.readouterr()
    assert cli.main(["--root", str(tmp_path), "run"]) == 0
    assert "지원방법이 바뀌어" not in capsys.readouterr().out  # 한 번만


def test_report_with_resume_items_and_wording(tmp_path):
    project = _project(tmp_path, resume=RESUME)
    text = (project.job_dir(KEY) / "report.md").read_text(encoding="utf-8")
    assert "위로 올릴 항목 (resume.yaml)" in text
    assert "[기술] 언어 — 필수 Python 기반 API 개발" in text
    # 공고는 'Spring Boot'·'Python'으로 쓰는데 이력서 글은 '스프링부트'·'파이썬'. '스프링부트' 안의 '스프링'은 따로 잡지 않는다
    assert "- 주문 API 개선: '스프링부트' → 'Spring Boot'" in text
    assert "- 언어: '파이썬' → 'Python'" in text
    assert "'스프링' →" not in text


def test_user_questions_trigger_rematch(tmp_path, capsys):
    project = _project(tmp_path)
    wb = load_workbook(project.jobs_xlsx)
    ws = wb[SHEET_QUESTIONS]
    ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "예시랩", 1, "어려운 문제를 해결한 경험을 적어주세요", 700
    wb.save(project.jobs_xlsx)
    capsys.readouterr()
    cli.main(["--root", str(tmp_path), "run"])
    out = capsys.readouterr().out
    assert "문항이(가) 바뀌어 매칭부터 다시" in out
    match = json.loads((project.job_dir(KEY) / "match.json").read_text(encoding="utf-8"))
    q = match["questions"][0]
    assert q["source"] == "직접입력" and q["type"] == "문제해결" and q["limit"] == 700
    assert q["cards"][0]["role"] == "주력"
    assert _job(project).status != "문항대기"  # 초안 단계(MVP-4) 앞에서 멈춘다


def test_changed_cards_trigger_rematch_and_resume_triggers_report(tmp_path, capsys):
    project = _project(tmp_path)
    project.experiences_file.write_text(CARDS.replace("[Python, PostgreSQL]", "[Python, PostgreSQL, MongoDB]"), encoding="utf-8")
    capsys.readouterr()
    cli.main(["--root", str(tmp_path), "run"])
    assert "경험 카드이(가) 바뀌어 매칭부터 다시" in capsys.readouterr().out
    project.resume_file.write_text(RESUME, encoding="utf-8")
    cli.main(["--root", str(tmp_path), "run"])
    assert "이력서 항목이 바뀌어 report부터 다시" in capsys.readouterr().out


def test_matching_rules_change_triggers_rematch(tmp_path, capsys):
    # 예전 버전이 만든 match.json(규칙 해시 없음)은 올린 뒤 첫 실행에서 다시 매칭한다
    project = _project(tmp_path)
    state = State.load(project.state_file)
    state.get(KEY).steps["match"].pop("rules_sha")
    state.save(project.state_file)
    capsys.readouterr()
    cli.main(["--root", str(tmp_path), "run"])
    assert "매칭 규칙이(가) 바뀌어 매칭부터 다시" in capsys.readouterr().out


def test_report_names_conditions_left_out_of_the_score(tmp_path):
    # 실측: '경력 3년 이상'이 "1개는 판단하지 않았습니다"에 묻혀 점수만 보였다
    analysis = {**ANALYSIS, "must": [*ANALYSIS["must"], "경력 3년 이상"], "nice": [*ANALYSIS["nice"], "정보처리기사 자격증"]}
    project = _project(tmp_path, analysis=analysis)
    text = (project.job_dir(KEY) / "report.md").read_text(encoding="utf-8")
    assert "**점수에 넣지 않은 조건** — 기술 낱말이 없어 경험 카드로 판단할 수 없습니다. 직접 판단하세요" in text
    assert "- 필수: 경력 3년 이상" in text
    assert "- 우대: 정보처리기사 자격증" in text


def test_example_resume_is_not_used_for_emphasis(tmp_path, capsys):
    shutil.copytree(Path(__file__).resolve().parents[1] / "me.example", tmp_path / "me.example")
    project = _project(tmp_path)  # init 이 me.example 을 me/ 로 복사 — resume.yaml 은 예시 그대로
    text = (project.job_dir(KEY) / "report.md").read_text(encoding="utf-8")
    assert "앞세울 경험 (내 이력서 항목이 없어 경험 카드 순위로)" in text
    assert "주문 조회 API 응답 시간 개선" not in text
    assert "me/resume.yaml 이 예시 그대로입니다" in text
    capsys.readouterr()
    cli.main(["--root", str(tmp_path), "next"])
    assert "me/resume.yaml 이 예시 그대로입니다" in capsys.readouterr().out.split(HEADER, 1)[1]


def test_example_resume_items_mixed_with_mine_are_dropped(tmp_path):
    shutil.copytree(Path(__file__).resolve().parents[1] / "me.example", tmp_path / "me.example")
    example = (tmp_path / "me.example" / "resume.yaml").read_text(encoding="utf-8")
    project = _project(tmp_path, resume=example.rstrip() + "\n" + RESUME.split("items:", 1)[1])
    text = (project.job_dir(KEY) / "report.md").read_text(encoding="utf-8")
    assert "위로 올릴 항목 (resume.yaml)" in text
    assert "[기술] 언어 — 필수 Python 기반 API 개발" in text
    assert "주문 조회 API 응답 시간 개선 —" not in text
    assert "me/resume.yaml 에 예시 항목 3개가 남아 있습니다(주문 조회 API 응답 시간 개선, 백엔드, 신규 팀원 온보딩 문서 정리)" in text


AI_CARDS = CARDS.replace("  verified: false\n", "  verified: false\n  source: sources/resume.pdf\n") + "".join(
    f"- id: extra{i}\n  title: 동아리 활동 {i}\n  tags: [협업]\n  summary: 활동\n  facts: []\n  usable_for: [협업]\n"
    "  verified: false\n  source: sources/resume.pdf\n"
    for i in range(3)
)


def test_unverified_notice_points_to_cards_in_use(tmp_path, capsys):
    # 실측: "AI가 뽑은 카드 31개가 확인 전" 이 매번 나와 31개를 당장 다 확인해야 하는 일로 읽혔다
    project = _project(tmp_path, cards=AI_CARDS)
    capsys.readouterr()
    cli.main(["--root", str(tmp_path), "next"])
    block = capsys.readouterr().out.split(HEADER, 1)[1]
    assert "공고 리포트·초안에 쓰이는 경험 카드 1개가 아직 확인 전입니다(api)" in block
    assert "나머지 3개는 쓰일 때 알려 드립니다" in block
    text = (project.job_dir(KEY) / "report.md").read_text(encoding="utf-8")
    assert "'동아리 활동 0'은(는) verified" not in text  # 점수 0 카드는 확인 필요에도 넣지 않는다


def test_dashboard_contents(tmp_path):
    project = _project(tmp_path)
    wb = load_workbook(project.dashboard_xlsx)
    ws = wb["검토"]
    header = [c.value for c in ws[1]]
    assert "갭" not in header and header[:5] == ["회사명", "직무", "결정", "D-day", "처리상태"]
    row = {h: ws.cell(row=2, column=i + 1) for i, h in enumerate(header)}
    assert row["회사명"].value == "예시랩"
    assert row["D-day"].value == "=DATE(2026,10,27)-TODAY()"
    assert isinstance(row["매칭점수"].value, int)
    assert row["문항"].value == "미입력" and row["초안"].value == "문항 없음"
    assert row["처리상태"].value == "문항대기"
    assert row["지원방법"].value == "홈페이지"
    assert row["결과"].hyperlink.target == "output/saramin-10000001/report.md"
    assert row["공고키"].value == KEY
    assert "입력 문제" not in wb.sheetnames


def test_dashboard_lists_sheet_errors_and_respects_lock(tmp_path, capsys):
    project = _project(tmp_path)
    wb = load_workbook(project.jobs_xlsx)
    wb["공고"]["A4"], wb["공고"]["C4"] = "링크이상", "ftp://x"
    wb.save(project.jobs_xlsx)
    cli.main(["--root", str(tmp_path), "dashboard"])
    assert "입력 문제" in load_workbook(project.dashboard_xlsx).sheetnames
    (tmp_path / "~$dashboard.xlsx").write_bytes(b"x")
    before = project.dashboard_xlsx.read_bytes()
    capsys.readouterr()
    assert cli.main(["--root", str(tmp_path), "dashboard"]) == 1
    assert "dashboard.xlsx가 열려 있습니다" in capsys.readouterr().out
    assert project.dashboard_xlsx.read_bytes() == before


def test_no_cards_still_gives_report(tmp_path):
    project = _project(tmp_path, cards=None)
    text = (project.job_dir(KEY) / "report.md").read_text(encoding="utf-8")
    assert "경험 카드가 없어 매칭하지 않았습니다" in text
    assert _job(project).status == "문항대기"
