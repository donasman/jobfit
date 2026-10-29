"""jobs.xlsx 읽기 — 헤더 이름으로 컬럼을 찾고, 행 단위로 검증한다.

문제 행은 건너뛰고 오류 목록에 남긴다. 전체 실행은 계속된다.
이 모듈은 파일을 절대 쓰지 않는다.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from jobfit.excel.workbook import (
    JOB_HEADERS,
    JOB_REQUIRED,
    QUESTION_HEADERS,
    QUESTION_REQUIRED,
    SHEET_JOBS,
    SHEET_QUESTIONS,
    SPACE_OPTIONS,
)
from jobfit.dates import parse_deadline
from jobfit.state import DECISION_DEFAULT, DECISIONS, TASKS

_WS = re.compile(r"[\s　 ]+")  # 공백·전각 공백·nbsp


@dataclass
class RowError:
    sheet: str
    row: int
    message: str

    def __str__(self) -> str:
        return f"[{self.sheet} {self.row}행] {self.message}"


@dataclass
class JobRow:
    row: int
    company: str
    title: str = ""
    url: str = ""
    deadline: str = ""  # YYYY-MM-DD
    deadline_note: str = ""  # 상시 등 날짜가 아닌 입력 원문
    memo: str = ""
    decision: str = DECISION_DEFAULT
    task: str = ""

    @property
    def display(self) -> str:
        return f"{self.company} / {self.title}" if self.title else self.company


@dataclass
class QuestionRow:
    row: int
    job_ref: str  # "회사명 / 직무" 또는 "회사명"
    number: int
    text: str
    limit: int | None = None
    count_spaces: bool = True


@dataclass
class ReadResult:
    jobs: list[JobRow] = field(default_factory=list)
    questions: list[QuestionRow] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors


def read_jobs_workbook(path: Path) -> ReadResult:
    result = ReadResult()
    # read_only=False: 열려 있는 파일도 읽는다. data_only=True: 수식 대신 값
    wb = load_workbook(path, data_only=True)
    if SHEET_JOBS not in wb.sheetnames:
        result.errors.append(RowError(SHEET_JOBS, 0, f"'{SHEET_JOBS}' 시트가 없습니다. 시트 이름을 되돌려 주세요."))
    else:
        _read_jobs_sheet(wb[SHEET_JOBS], result)
    if SHEET_QUESTIONS in wb.sheetnames:
        _read_questions_sheet(wb[SHEET_QUESTIONS], result)
    return result


# --- 시트별 ---


def _read_jobs_sheet(ws, result: ReadResult) -> None:
    cols = _header_map(ws, JOB_HEADERS, SHEET_JOBS, result)
    if cols is None:
        return
    for row in ws.iter_rows(min_row=2):
        r = row[0].row
        get = lambda name: _cell(row, cols.get(name))  # noqa: E731
        company = _clean(get("회사명"))
        title = _clean(get("직무"))
        url = _link(row, cols.get("링크"))
        memo = _clean(get("메모"))
        if not any((company, title, url, memo)):
            continue  # 빈 행
        if not company:
            result.errors.append(RowError(SHEET_JOBS, r, "회사명이 비어 있습니다"))
            continue
        if url and not re.match(r"^https?://", url, re.I):
            result.errors.append(RowError(SHEET_JOBS, r, f"링크가 http(s)://로 시작하지 않습니다: {url[:40]}"))
            continue
        deadline, note, err = parse_deadline(get("마감일"))
        if err:
            result.errors.append(RowError(SHEET_JOBS, r, err))
            continue
        decision = _clean(get("결정")) or DECISION_DEFAULT
        if decision not in DECISIONS:
            result.errors.append(RowError(SHEET_JOBS, r, f"결정 값이 목록에 없습니다: {decision}"))
            continue
        task = _clean(get("작업"))
        if task not in TASKS:
            result.errors.append(RowError(SHEET_JOBS, r, f"작업 값이 목록에 없습니다: {task}"))
            continue
        result.jobs.append(
            JobRow(row=r, company=company, title=title, url=url, deadline=deadline, deadline_note=note, memo=memo, decision=decision, task=task)
        )


def _read_questions_sheet(ws, result: ReadResult) -> None:
    cols = _header_map(ws, QUESTION_HEADERS, SHEET_QUESTIONS, result)
    if cols is None:
        return
    for row in ws.iter_rows(min_row=2):
        r = row[0].row
        get = lambda name: _cell(row, cols.get(name))  # noqa: E731
        job_ref = _clean(get("공고"))
        text = _clean(get("문항"), keep_newlines=True)
        if not job_ref and not text:
            continue
        if not job_ref:
            result.errors.append(RowError(SHEET_QUESTIONS, r, "공고가 비어 있습니다"))
            continue
        if not text:
            result.errors.append(RowError(SHEET_QUESTIONS, r, "문항이 비어 있습니다"))
            continue
        number_raw = get("번호")
        try:
            number = int(number_raw) if number_raw not in (None, "") else 0
        except (TypeError, ValueError):
            result.errors.append(RowError(SHEET_QUESTIONS, r, f"번호는 숫자여야 합니다: {number_raw}"))
            continue
        limit_raw = get("글자수")
        limit: int | None
        if limit_raw in (None, ""):
            limit = None
        else:
            m = re.search(r"\d+", str(limit_raw).replace(",", ""))
            if not m:
                result.errors.append(RowError(SHEET_QUESTIONS, r, f"글자수는 숫자만 적습니다: {limit_raw}"))
                continue
            limit = int(m.group())
        spaces = _clean(get("공백포함")) or SPACE_OPTIONS[0]
        if spaces not in SPACE_OPTIONS:
            result.errors.append(RowError(SHEET_QUESTIONS, r, f"공백포함은 포함/제외 중 하나입니다: {spaces}"))
            continue
        result.questions.append(QuestionRow(row=r, job_ref=job_ref, number=number, text=text, limit=limit, count_spaces=(spaces == "포함")))


# --- 공통 ---


def _header_map(ws, expected: tuple[str, ...], sheet: str, result: ReadResult) -> dict[str, int] | None:
    """1행에서 헤더 이름 -> 열 인덱스(0부터). 필수 헤더가 없으면 None."""
    header_row = next(ws.iter_rows(min_row=1, max_row=1), ())
    cols: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        name = _clean(cell.value)
        if name and name not in cols:
            cols[name] = idx
    required = JOB_REQUIRED if sheet == SHEET_JOBS else QUESTION_REQUIRED
    missing = [h for h in expected if h not in cols]
    missing_required = [h for h in required if h not in cols]
    if missing_required:
        result.errors.append(RowError(sheet, 1, f"필수 헤더가 없습니다: {', '.join(missing_required)}. 헤더 이름을 되돌려 주세요."))
        return None
    for h in missing:
        result.errors.append(RowError(sheet, 1, f"헤더가 없어 그 열은 비어 있는 것으로 봅니다: '{h}'"))
    return cols


def _cell(row, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    return row[idx].value


def _link(row, idx: int | None) -> str:
    """링크 칸. 엑셀 하이퍼링크에 표시 글자를 따로 넣었으면 주소 쪽을 쓴다."""
    if idx is None or idx >= len(row):
        return ""
    cell = row[idx]
    value = _clean(cell.value)
    target = getattr(getattr(cell, "hyperlink", None), "target", None)
    if target and not re.match(r"^https?://", value, re.I):
        return str(target).strip()
    return value


def _clean(value, keep_newlines: bool = False) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value).strip()
    if keep_newlines:
        return "\n".join(_WS.sub(" ", line).strip() for line in s.splitlines()).strip()
    return _WS.sub(" ", s).strip()
