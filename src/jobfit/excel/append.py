"""jobs.xlsx에 공고 한 줄 추가 — `jobfit add` 전용. 스크립트가 jobs.xlsx에 쓰는 유일한 곳이다.

사용자가 명시적으로 부를 때만 쓴다. 엑셀이 열려 있으면 쓰지 않는다. 쓰기 전에 백업을 남긴다.
공고 시트의 비어 있는 첫 행에, 사용자 열 값만 쓴다. 결정·작업은 비워 둔다.
"""

from __future__ import annotations

import os
import shutil
import tempfile
import warnings
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from jobfit.excel.workbook import JOB_DISPLAY_HEADER, JOB_HEADERS, MAX_ROWS, SHEET_JOBS, display_formula, ensure_validations, header_letters

BACKUP_KEEP = 10


class ExcelBusy(Exception):
    """엑셀(또는 LibreOffice)이 파일을 열고 있어 쓸 수 없다."""


def lock_files(path: Path) -> list[Path]:
    """편집 프로그램이 파일을 열 때 옆에 만드는 잠금 파일들."""
    return [path.with_name("~$" + path.name), path.with_name(f".~lock.{path.name}#")]


def append_job_row(path: Path, backup_dir: Path, company: str, title: str = "", url: str = "", deadline: str = "", memo: str = "") -> int:
    """추가한 행 번호를 돌려준다."""
    if any(p.exists() for p in lock_files(path)):
        raise ExcelBusy("jobs.xlsx가 열려 있습니다. 엑셀을 닫고 다시 실행하세요")
    backup(path, backup_dir)
    with warnings.catch_warnings():
        # 엑셀이 저장한 확장 형식(목록 검사 등)을 openpyxl이 버린다는 경고 — ensure_validations가 다시 단다
        warnings.simplefilter("ignore", UserWarning)
        wb = load_workbook(path)
    ws = wb[SHEET_JOBS]
    cols = header_letters(ws)
    if "회사명" not in cols:
        raise ValueError("공고 시트에 '회사명' 헤더가 없습니다. 헤더 이름을 되돌려 주세요")
    row = _first_empty_row(ws, cols)
    values = {
        "회사명": company,
        "직무": title,
        "링크": url,
        "마감일": date.fromisoformat(deadline) if deadline else "",
        "메모": memo,
    }
    for name, value in values.items():
        if name in cols and value not in ("", None):
            cell = ws[f"{cols[name]}{row}"]
            cell.value = value
            if name == "마감일":
                cell.number_format = "yyyy-mm-dd"
    if JOB_DISPLAY_HEADER in cols and row > MAX_ROWS:
        ws[f"{cols[JOB_DISPLAY_HEADER]}{row}"] = display_formula(cols["회사명"], cols.get("직무"), row)
    ensure_validations(wb)
    _save_atomic(wb, path)
    return row


def backup(path: Path, backup_dir: Path) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    dst = backup_dir / f"jobs_{stamp}.xlsx"
    n = 1
    while dst.exists():
        dst = backup_dir / f"jobs_{stamp}_{n}.xlsx"
        n += 1
    shutil.copy2(path, dst)
    for old in sorted(backup_dir.glob("jobs_*.xlsx"))[:-BACKUP_KEEP]:
        old.unlink(missing_ok=True)
    return dst


def _first_empty_row(ws, cols: dict[str, str]) -> int:
    """사용자 열이 전부 빈 첫 행. 공고표시 수식 열은 보지 않는다(500행까지 차 있다)."""
    letters = [cols[h] for h in JOB_HEADERS if h in cols]
    for r in range(2, ws.max_row + 2):
        if all(ws[f"{L}{r}"].value in (None, "") for L in letters):
            return r
    return ws.max_row + 1


def _save_atomic(wb, path: Path) -> None:
    fd, tmp = tempfile.mkstemp(prefix="~jobfit-", suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError as e:
        raise ExcelBusy("jobs.xlsx에 쓸 수 없습니다(열려 있나요?). 엑셀을 닫고 다시 실행하세요") from e
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)
