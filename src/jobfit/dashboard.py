"""dashboard.xlsx — 스크립트 전용 출력. state.json과 output/에서 매번 통째로 다시 만든다.

사용자가 이 파일을 고쳐도 다음 실행에서 사라진다(안내 줄에 적혀 있다). 입력은 jobs.xlsx에만.
"""

from __future__ import annotations

import os
import tempfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from jobfit.drafts import lint_summary, load_lint
from jobfit.excel.append import lock_files
from jobfit.match import load_analysis, load_match
from jobfit.paths import Project
from jobfit.report import FILE_REPORT
from jobfit.state import DECISIONS_CLOSED, DECISIONS_PAUSED, STATUS_AI, STATUS_NEED_QUESTIONS, STATUS_REVIEW, Job, State

SHEET_REVIEW = "검토"
SHEET_INPUT = "입력 문제"
HEADERS = ("회사명", "직무", "결정", "D-day", "처리상태", "매칭점수", "문항", "초안", "다음 할 일", "오류", "결과", "공고키")
WIDTHS = (16, 26, 9, 8, 10, 9, 16, 12, 40, 30, 14, 24)

_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
_SOON_FILL = PatternFill("solid", fgColor="FFF2CC")
_GREY = Font(color="999999")


class DashboardBusy(Exception):
    """dashboard.xlsx가 열려 있어 쓸 수 없다."""


def build_dashboard(project: Project, state: State) -> Path:
    path = project.dashboard_xlsx
    if any(p.exists() for p in lock_files(path)):
        raise DashboardBusy("dashboard.xlsx가 열려 있습니다. 닫고 다시  jobfit dashboard")
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_REVIEW
    ws.append(HEADERS)
    for col, width in enumerate(WIDTHS, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    jobs = sorted(state.jobs.values(), key=_order)
    for r, job in enumerate(jobs, start=2):
        values = _row(project, job)
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value)
        if job.deadline:
            y, m, d = (int(x) for x in job.deadline.split("-"))
            cell = ws.cell(row=r, column=HEADERS.index("D-day") + 1, value=f"=DATE({y},{m},{d})-TODAY()")
            cell.number_format = "0"
        report = project.job_dir(job.key) / FILE_REPORT
        if report.exists():
            cell = ws.cell(row=r, column=HEADERS.index("결과") + 1, value="report.md")
            cell.hyperlink = project.rel(report)
            cell.style = "Hyperlink"
        if not job.is_active:
            for c in range(1, len(HEADERS) + 1):
                ws.cell(row=r, column=c).font = _GREY
    last = max(len(jobs) + 1, 2)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last}"
    d = get_column_letter(HEADERS.index("D-day") + 1)
    ws.conditional_formatting.add(f"{d}2:{d}{last}", FormulaRule(formula=[f"AND(ISNUMBER({d}2),{d}2>=0,{d}2<=7)"], fill=_SOON_FILL))
    ws.conditional_formatting.add(f"{d}2:{d}{last}", FormulaRule(formula=[f"AND(ISNUMBER({d}2),{d}2<0)"], font=_GREY))
    note_row = last + 2
    ws.cell(row=note_row, column=1, value=f"스크립트가 매번 다시 만드는 파일입니다. 여기 고친 내용은 사라집니다 — 입력은 jobs.xlsx에. (만든 날 {date.today().isoformat()})").font = _GREY

    if state.sheet_errors:
        wi = wb.create_sheet(SHEET_INPUT)
        wi.append(("jobs.xlsx에서 읽지 못한 행",))
        wi["A1"].font = Font(bold=True)
        wi.column_dimensions["A"].width = 100
        for e in state.sheet_errors:
            wi.append((e,))

    _save_atomic(wb, path)
    return path


def _row(project: Project, job: Job) -> list:
    match = load_match(project, job)
    analysis = load_analysis(project, job)
    questions = match.get("questions")
    if questions is None:
        source = "직접입력" if job.user_questions else "공고추출"
        n = len(job.user_questions) if job.user_questions else len(analysis.get("questions_found", [])) if analysis else None
    else:
        n = len(questions)
        source = questions[0]["source"] if questions else ""
    if n is None:
        q_text = ""
    elif n == 0:
        q_text = "미입력"
    else:
        q_text = f"{n}개 ({source})"
    if job.step_done("draft"):
        draft = lint_summary(load_lint(project, job))
    else:
        draft = "문항 없음" if job.step_done("report") and n == 0 else ""
    score = match.get("score") if match else None
    return [
        job.company,
        job.title,
        job.decision,
        "",
        job.status,
        score if score is not None else "",
        q_text,
        draft,
        _next(job),
        job.error,
        "",
        job.key,
    ]


def _next(job: Job) -> str:
    if job.decision in DECISIONS_CLOSED:
        return ""
    if job.decision in DECISIONS_PAUSED:
        return "처리하지 않음 (결정을 미정으로 돌리면 다시 처리)"
    if job.task_seen:
        return "작업 칸을 비워 주세요"
    if job.status == STATUS_AI:
        return job.next_hint or "AI 프롬프트 복사 (메뉴 3)"
    if job.status == STATUS_NEED_QUESTIONS:
        return "문항 시트에 문항 입력, 또는 결정 열에서 지원 여부 결정"
    if job.status == STATUS_REVIEW:
        return "report.md 확인 후 결정 열에서 지원 여부 결정"
    return job.next_hint


def _order(job: Job):
    group = 0 if job.is_active else (1 if job.is_paused else 2)
    return (group, job.deadline or "9999-99-99", job.row)


def _save_atomic(wb, path: Path) -> None:
    fd, tmp = tempfile.mkstemp(prefix="~jobfit-", suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError as e:
        raise DashboardBusy("dashboard.xlsx에 쓸 수 없습니다(열려 있나요?). 닫고 다시  jobfit dashboard") from e
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)
